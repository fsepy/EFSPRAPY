from typing import Dict

import openpyxl


def xlsx_to_dict(fp: str) -> Dict[str, dict]:
    # Load the workbook
    wb = openpyxl.load_workbook(fp)

    # Select the first worksheet
    sheet = wb.active

    # Get the headers (first row) and the row headers (first column)
    headers = [cell.value for cell in sheet[1]]
    headers = headers[1:]
    row_headers = [cell.value for cell in sheet['A'][1:]]

    # Iterate through the sheet and create the nested dictionary
    result = {}
    for header, col_idx in zip(headers, range(2, len(headers) + 2)):
        col_values = {}
        for row_header, row_idx in zip(row_headers, range(2, len(row_headers) + 2)):
            col_values[row_header] = sheet.cell(row=row_idx, column=col_idx).value
        result[header] = col_values

    return result


def dict_to_xlsx(data: Dict[str, dict], fp: str):
    """

    :param data:
    :param fp:
    :return:

    data = {
        "Column1": {"Row1": "A1", "Row2": "A2", "Row3": "A3"},
        "Column2": {"Row1": "B1", "Row2": "B2", "Row3": "B3"},
        "Column3": {"Row1": "C1", "Row2": "C2", "Row3": "C3"},
    }
    """
    # check all nested dict
    row_headers = None
    for k, v in data.items():
        if row_headers:
            assert row_headers == v.keys()
        row_headers = v.keys()

    # Create a new workbook and add a worksheet
    wb = openpyxl.Workbook()
    ws = wb.active

    # Write the column headers
    for col, column_name in enumerate(data.keys(), start=1):
        ws.cell(row=1, column=col + 1, value=column_name)

    # Write the row headers
    for row, row_name in enumerate(row_headers):
        ws.cell(row=row + 2, column=1, value=row_name)

    # Write the data to the worksheet
    # Iterate through columns
    for col_index, (column_name, data_) in enumerate(data.items()):
        # The actual column number in Excel (starts from B=2)
        excel_col = col_index + 2  # +1 for 0-based enumerate, +1 for skipping column A

        # Iterate through the *standardized* row headers
        for row_index, row_name in enumerate(row_headers):
            # The actual row number in Excel (starts from row 2)
            excel_row = row_index + 2  # +1 for 0-based enumerate, +1 for skipping row 1

            # Look up the value using the row_name from the standardized list
            # This ensures the correct value is placed in the correct row
            value = data_.get(row_name, "")
            # Use .get() for safety in case of missing keys (though assert should prevent this)

            ws.cell(row=excel_row, column=excel_col, value=value)

    # Save the workbook to an XLSX file
    wb.save(fp)
